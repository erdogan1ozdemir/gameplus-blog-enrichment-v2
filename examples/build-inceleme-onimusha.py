# -*- coding: utf-8 -*-
"""Onimusha: Way of the Sword incelemesi -> v10.26 enriched HTML + doc (metin + HTML).

Brief: Game+ Blogları/Game+ Blog Briefleri.xlsx (satir: Onimusha: Way of the Sword).
Tarz: references/inceleme-yazisi.md (sayilar asagi yuvarlanir, basin adi yok, GFN tek paragraf,
sistem gereksinimleri tablosu yok, yas kisitli video yok).

Kaynaklar (24 Eylul 2026'da dogrulandi, her sayfada oyun adi h1/title ile kontrol edildi):
- Steam appdetails (appid 2638890): Capcom, tek kisilik, HDR, tam kontrolcu, Turkce YOK,
  sistem gereksinimleri (min GTX 1660 / RX 5500 XT, onerilen RTX 2060 Super / RX 6600, 50 GB).
- Steam appreviews: tum diller 23.850 inceleme, 21.578 olumlu (%90, Cok Olumlu);
  Turkce 178 inceleme, 163 olumlu (%91).
- OpenCritic (game/20764): Top Critic Average 86, %95 tavsiye, 159 elestirmen, puan araligi 4/5 - 10/10.
- Metacritic: 85 (101 elestirmen), kullanici 8.6 (836 oy).
- HowLongToBeat (game/160598): ana 22 saat, ana+ekstra 29,5 saat, tamamlama 49 saat.
- Wikipedia: 4 Eylul 2026 cikis, platformlar, RE Engine, Dawn of Dreams'ten (2006) sonra ilk ana oyun,
  ilk gun 1 milyon satis, Musashi'nin Toshiro Mifune'den esinlenmesi, Oni Gauntlet ruh renkleri,
  Issen/Break Issen, yonetmen Satoru Nihei "Soulslike degil", hikaye onceki oyunlardan bagimsiz.
- GAME+ 3 Eylul 2026 GFN Thursday yazisi: oyun cikis gunu GeForce NOW'da (headline ile dogrulandi).
- Gomulu video: Capcom USA Launch Trailer (Gbmd6YFm5oU), tarayicida yas kisiti testi gecti.
"""
import os
import re
import sys

SKILL = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "scripts")
sys.path.insert(0, SKILL)
from gameplus_blog_components import *  # noqa: F401,F403

OUT = os.environ.get("GP_OUT", ".")
TITLE = "Onimusha: Way of the Sword İncelemesi: Seri 20 Yıl Sonra Döndü"
SLUG = "onimusha-way-of-the-sword-incelemesi"
AD = "Onimusha: Way of the Sword"

P = []


def h1(t): P.append(("h1", t))
def h2(t): P.append(("h2", t))
def p(t): P.append(("p", t))
def yt(vid, baslik): P.append(("yt", (vid, baslik)))


h1(TITLE)

OC = "https://opencritic.com/game/20764/onimusha-way-of-the-sword"
MC = "https://www.metacritic.com/game/onimusha-way-of-the-sword/"
HLTB = "https://howlongtobeat.com/game/160598"
GFN_YAZI = "https://gameplus.com.tr/blog/gfn-thursday-geforce-now-da-bu-hafta-3-eylul-2026"

p("Edo döneminin başında Kyoto, Genma adı verilen iblislerin istilası altında. Şehrin sokaklarında "
  "kılıcını çeken isim ise Japonya'nın en ünlü kılıç ustalarından biri: Miyamoto Musashi. "
  f"<strong>{AD}</strong>, bu kurulumla Capcom'un uzun süredir uykuda olan serisini geri getiriyor.")

p("Oyun 4 Eylül 2026'da çıktı ve serinin 2006 tarihli Onimusha: Dawn of Dreams'ten bu yana ilk ana "
  "oyunu oldu. Capcom'un açıklamasına göre ilk gün 1 milyon satışa ulaştı. Peki 20 yıllık "
  "aradan sonra gelen bu dönüş, beklentiyi karşılıyor mu?")

h2(f"{AD} Nedir?")
p(f"{AD}, Capcom'un geliştirip yayınladığı, Japon tarihi ile karanlık fanteziyi birleştiren bir "
  "aksiyon oyunu. RE Engine ile geliştirilen oyun PC, PlayStation 5, Xbox Series X|S ve Nintendo "
  "Switch 2 için çıktı. Serinin tanıdık formülü korunuyor: kılıç dövüşü, iblislerden toplanan "
  "ruhlar ve bu ruhlarla güçlenen bir kahraman.")
yt("Gbmd6YFm5oU", "Onimusha: Way of the Sword Çıkış Fragmanı")
p("Bu bir remake değil. İlk oyunun yeniden düzenlenmiş sürümü 2019'da çıkmıştı; Way of the Sword ise "
  "sıfırdan yazılmış yeni bir hikaye anlatıyor. Proje 2020'de onay aldı, yapımcı koltuğunda Akihito "
  "Kadowaki, yönetmen koltuğunda Satoru Nihei oturuyor. Hikaye önceki oyunlardan ve Netflix'teki "
  "animasyon dizisinden bağımsız ilerlediği için seriye ilk kez başlayacaklar için de uygun bir "
  "giriş noktası.")

h2(f"{AD} Hikayesi Neyi Anlatıyor?")
p("Oyun, Genma istilasıyla karanlığa gömülen Kyoto'da, iblis gücü taşıyan bir eldivenle Genma'ya "
  "karşı savaşan Miyamoto Musashi'yi anlatıyor. Oni Gauntlet adı verilen bu eldiven, öldürülen "
  "iblislerin ruhlarını emiyor ve Musashi'ye insanüstü bir güç kazandırıyor.")
p("Musashi'nin görünüşü ve duruşu, Japon sinemasının efsane oyuncusu Toshiro Mifune'den "
  "esinleniyor. Yolculuk boyunca Ono no Takamura, Izumo no Okuni ve Sasaki Ganryu gibi tarihten "
  "ve efsanelerden tanınan isimlerle karşılaşıyorsun. Hikaye ciddi ve yer yer sert bir tonda "
  "ilerliyor; buna karşılık anlatımın temposu bölümden bölüme değişiyor ve bazı ara bölümler "
  "hikayeyi yavaşlatıyor.")

h2(f"{AD}'da Dövüş Nasıl İşliyor?")
p("Dövüş, zamanlamaya dayalı kılıç düellolarının üzerine kurulu: düşmanın saldırısını doğru anda "
  "savuşturmak, açığı yakalamak ve tek hamlede bitirmek oyunun ana döngüsü. Temel parçalar "
  "şöyle:")
P.append(("liste", [
    "<strong>Savuşturma ve saptırma:</strong> saldırıyı tam zamanında karşılamak düşmanın dengesini "
    "bozuyor; blok tutmak ise yalnız zaman kazandırıyor.",
    "<strong>Issen ve Break Issen:</strong> serinin imzası olan tek vuruşluk karşı saldırı geri "
    "dönüyor; Break Issen ile bu sistem denge kırma mekaniğine bağlanıyor.",
    "<strong>Renkli ruhlar:</strong> Oni Gauntlet'in topladığı sarı ruhlar can yeniliyor, kırmızı "
    "ruhlar gelişim için harcanıyor, mavi ruhlar iblis silahlarının gücünü dolduruyor.",
    "<strong>Çevre:</strong> bir masayı devirmek gibi çevre etkileşimleri dövüşte taktik seçenek "
    "olarak kullanılabiliyor.",
]))
p("Capcom, kılıç hareketlerini gerçek kılıç ustalarıyla hareket yakalama çekimleri yaparak "
  "hazırladı ve bu emek dövüşün hissine yansıyor. Eleştirmenlerin büyük bölümü dövüş sistemini ve "
  "boss mücadelelerini oyunun en güçlü yanı olarak görüyor; oyuncu yorumlarında da en sık övülen "
  "başlık yine dövüş.")

h2(f"{AD} Sekiro Gibi mi, Zor mu?")
p(f"Hayır, {AD} tam olarak Sekiro gibi değil: savuşturmayı merkeze alması benzer, ancak yönetmen "
  "Satoru Nihei oyunun bir Soulslike olmadığını açıkça söylüyor. Ölüm cezası, kontrol noktası "
  "düzeni ve genel zorluk, türün sert örneklerine göre daha erişilebilir tutulmuş.")
p("Bu tercih iki farklı tepki topluyor. Sekiro ya da Nioh gibi oyunlardan gelen oyuncuların bir "
  "kısmı oyunu fazla kolay buluyor ve yol arkadaşlarının sık sık ipucu vermesini gereksiz "
  "görüyor. Türe yeni başlayanlar içinse bu erişilebilirlik, Onimusha'yı savuşturma odaklı "
  "dövüşe girmek için rahat bir başlangıç noktası yapıyor.")

h2(f"{AD} Puanları: Metacritic ve OpenCritic")
p(f"{AD}, 24 Eylül 2026 itibarıyla <a href=\"{OC}\">OpenCritic</a>'te 86, "
  f"<a href=\"{MC}\">Metacritic</a>'te 85 puana sahip; OpenCritic'teki eleştirmenlerin %95'i oyunu "
  "tavsiye ediyor. Tablo aynı tarihteki değerleri gösteriyor:")
P.append(("puan_tablo", None))
p("Eleştirmen notları 4/5 ile 10/10 arasında değişiyor ve dağılım belirgin biçimde olumlu tarafta "
  "toplanıyor. En yüksek notları verenler dövüşü, boss tasarımını ve karanlık Kyoto atmosferini "
  "öne çıkarıyor; daha mesafeli duranlar yan içeriklerin dolgu hissi vermesini, temponun yer yer "
  "düşmesini ve zorluğun düşük kalmasını gerekçe gösteriyor.")

h2(f"Oyuncular {AD} Hakkında Ne Diyor?")
p("Oyuncu tarafı da eleştirmenlerle aynı yönde: Steam'de 23 binden fazla incelemenin %90'ı olumlu "
  "ve oyun \"Çok Olumlu\" etiketi taşıyor. Yerli oyuncu topluluklarının Steam'de yazdığı 170'ten "
  "fazla incelemenin de %91'i olumlu.")
p("Globaldeki oyuncu toplulukları en çok dövüşün akıcılığını, boss savaşlarını ve PC tarafındaki "
  "optimizasyonu övüyor. Eleştiriler ise üç başlıkta toplanıyor: bazı açık alanların boş kalması ve "
  "yan görevlerin tekrar hissi vermesi, oyunun zorluk seviyesinin düşük bulunması ve yol "
  "arkadaşlarının oyuncuyu fazla yönlendirmesi. Yerli oyuncu topluluklarında bunlara bir madde "
  "daha ekleniyor: Türkçe dil desteğinin olmaması.")

h2(f"{AD} Türkçe mi, Kaç Saat Sürüyor?")
p(f"Hayır, {AD} Türkçe dil desteği sunmuyor; 24 Eylül 2026 itibarıyla Steam sayfasındaki dil "
  "listesinde Türkçe yer almıyor. \"Türkçe yama\" aramalarının sık yapılması da bu yüzden; resmi "
  "bir Türkçe güncellemesi duyurulmuş değil.")
p(f"Süre tarafında oyun orta uzunlukta. <a href=\"{HLTB}\">HowLongToBeat</a> verilerine göre ana "
  "hikaye 20 saatin biraz üzerinde, yan içeriklerle birlikte 30 saat civarı sürüyor; her şeyi "
  "tamamlamak isteyenler 50 saate yakın zaman ayırıyor. Donanım tarafında minimum listede GTX 1660 "
  "seviyesi, önerilen listede RTX 2060 Super sınıfı bir ekran kartı görünüyor; kurulum için 50 GB "
  "alan isteniyor.")

h2(f"{AD} Alınır mı?")
p("Kılıç dövüşünü seviyorsan ve ağır bir Soulslike yerine akıcı bir aksiyon oyunu arıyorsan, "
  f"{AD} güçlü bir seçim. Şunları arıyorsan oyun tam sana göre:")
P.append(("liste2", [
    "<strong>Savuşturma odaklı dövüş sevenler:</strong> zamanlamaya dayalı düellolar ve Issen karşı "
    "saldırısı oyunun en çok övülen tarafı.",
    "<strong>Eski Onimusha hayranları:</strong> ruh toplama, iblis silahları ve Genma'ya karşı "
    "savaş gibi serinin temel unsurları yerinde.",
    "<strong>Japon tarihi ve samuray temasını sevenler:</strong> Edo dönemi Kyoto'su ve tarihi "
    "karakterler güçlü bir atmosfer kuruyor.",
]))
p("Beklentini ayarlaman gereken noktalar ise şöyle:")
P.append(("liste", [
    "<strong>Zorlu bir meydan okuma arıyorsan:</strong> oyun türün sert örneklerine göre daha "
    "erişilebilir ve kolay bulunabilir.",
    "<strong>Türkçe dil desteği bekliyorsan:</strong> oyunda resmi Türkçe arayüz ya da altyazı yok.",
    "<strong>Dolu bir açık dünya bekliyorsan:</strong> yan içeriklerin bir kısmı tekrar hissi "
    "verebiliyor.",
]))

h2("Genel Değerlendirme")
p(f"{AD}, serinin 20 yıllık aradan sonra geri dönüşünü sağlam bir temele oturtuyor. Zamanlamaya "
  "dayalı kılıç dövüşü, boss savaşları ve Genma istilasındaki Kyoto'nun atmosferi, oyunu hem eski "
  "hayranlar hem de seriye yeni başlayanlar için çekici kılıyor.")
p("Eksikleri de görünüyor: yan içerikler yer yer dolgu hissi veriyor, tempo bazı bölümlerde "
  "düşüyor ve zorluk seviyesi deneyimli oyuncular için düşük kalabiliyor. Yine de 86 ortalama ve "
  "eleştirmenlerin %95'inin tavsiye etmesi, oyunun bu eksiklere rağmen serinin güçlü bir dönüşü "
  "olduğunu gösteriyor. Benzer tarzda başka yapımlar arıyorsan GeForce NOW kütüphanesindeki aksiyon "
  "oyunları iyi bir başlangıç noktası.")

h2(f"{AD} GeForce NOW'da Oynanır mı?")
p(f"Evet, {AD} 4 Eylül'deki çıkışıyla aynı gün GeForce NOW kütüphanesine eklendi; oyunun buluta "
  f"gelişi <a href=\"{GFN_YAZI}\">3 Eylül 2026 tarihli GFN Thursday yazısında</a> da yer aldı. 50 "
  "GB'lık kurulumu beklemeden buluttan başlatabiliyor, önerilen listedeki ekran kartına sahip "
  "olmasan da farklı cihazlardan oynayabiliyorsun; işi bulut tarafındaki RTX sunucuları yapıyor ve "
  "kayıtların mağaza hesabında durduğu için kaldığın yerden devam ediyorsun.")

h2("Sıkça Sorulan Sorular")

FAQ = [
    (f"{AD} Türkçe dil desteği var mı?",
     "Hayır. Oyunda resmi Türkçe arayüz, altyazı ya da seslendirme bulunmuyor ve 24 Eylül 2026 "
     "itibarıyla Capcom tarafından duyurulmuş bir Türkçe güncellemesi yok. Oyunu İngilizce arayüz ve "
     "altyazıyla oynaman gerekiyor."),
    (f"{AD} kaç saatte bitiyor?",
     "HowLongToBeat verilerine göre ana hikaye 20 saatin biraz üzerinde bitiyor. Yan içeriklerle "
     "birlikte süre 30 saat civarına çıkıyor, her şeyi tamamlamak isteyenler ise 50 saate yakın "
     "zaman ayırıyor."),
    (f"{AD} ne zaman çıktı, hangi platformlarda var?",
     "Oyun 4 Eylül 2026'da PC, PlayStation 5, Xbox Series X|S ve Nintendo Switch 2 için çıktı. PC'de "
     "Steam üzerinden satın alınabiliyor; Capcom'un açıklamasına göre ilk gün 1 milyon satışa ulaştı."),
    (f"{AD} remake mi?",
     "Hayır. Way of the Sword serinin yeni bir ana oyunu ve 2006 tarihli Onimusha: Dawn of Dreams'ten "
     "sonra gelen ilk ana oyun. İlk Onimusha'nın yeniden düzenlenmiş sürümü ise 2019'da ayrı olarak "
     "yayınlanmıştı."),
    (f"{AD} Sekiro'dan zor mu?",
     "Hayır, genel olarak daha erişilebilir. Savuşturma odaklı dövüşüyle Sekiro'yu hatırlatıyor ama "
     "yönetmen Satoru Nihei oyunun Soulslike olmadığını söylüyor. Deneyimli oyuncuların bir kısmı "
     "zorluk seviyesini düşük buluyor."),
    ("Önceki Onimusha oyunlarını oynamadan Way of the Sword oynanır mı?",
     "Oynanır. Hikaye önceki oyunlardan ve Netflix'teki animasyon dizisinden bağımsız ilerliyor. Seriye "
     "ilk kez başlayanlar için Genma, ruh toplama ve Oni Gauntlet gibi unsurlar oyunun içinde baştan "
     "anlatılıyor."),
    (f"{AD} sistem gereksinimleri nedir?",
     "Minimum listede Windows 11, Intel Core i5-8400 ya da AMD Ryzen 3 3100 işlemci, 16 GB RAM ve GTX "
     "1660 ya da RX 5500 XT ekran kartı isteniyor. Önerilen listede RTX 2060 Super ya da RX 6600 yer "
     "alıyor; kurulum için 50 GB alan gerekiyor."),
    (f"{AD} GeForce NOW'da oynanabiliyor mu?",
     "Evet. Oyun çıkış günü GeForce NOW kütüphanesine eklendi ve bulut üzerinden oynanabiliyor. "
     "Oynayabilmen için oyuna Steam gibi desteklenen bir mağazada sahip olman gerekiyor; GeForce NOW "
     "oyunu satmaz, kütüphanendeki oyunu bulutta çalıştırır."),
]

PUAN_SATIRLARI = [
    ["OpenCritic", "86 (Top Critic Average)", "159 eleştirmen, %95 tavsiye ediyor"],
    ["Metacritic", "85", "101 eleştirmen"],
    ["Metacritic kullanıcı puanı", "8.6", "800'den fazla oy"],
    ["Eleştirmen puan aralığı", "4/5 ile 10/10 arası", "En düşük ve en yüksek notlar"],
    ["Steam kullanıcı incelemeleri", "%90 olumlu", "23 binden fazla inceleme, Çok Olumlu"],
    ["Steam Türkçe yorumlar", "%91 olumlu", "170'ten fazla inceleme"],
]
puan_tablo = render_table(["Kaynak", "Puan", "Kapsam"], PUAN_SATIRLARI, first_col_strong=False,
                          title=f"{AD} Puan Tablosu (24 Eylül 2026)")

editor_note = render_editor_note(
    "Seriye yeniysen önceki oyunları oynamadan başlayabilirsin; hikaye kendi başına duruyor. "
    "Savuşturma zamanlamasına alıştıkça dövüş hızla akmaya başlıyor, ilk saatlerde çok takılmadan "
    "ilerleyebilirsin.")

hatirlatma = render_highlight(
    f"GeForce NOW oyun satmaz, sahip olduğun oyunları bulutta çalıştırır. {AD}'u buluttan "
    "açabilmen için oyuna Steam veya desteklenen diğer mağazalardan birinde sahip olman gerekir.")

cta_paketler = render_cta_paketler(
    "Kyoto'ya 50 GB indirme beklemeden gir",
    "GeForce NOW, kütüphanendeki desteklenen oyunları güçlü bir ekran kartına ihtiyaç duymadan bulut "
    "üzerinden çalıştırır. Paketleri karşılaştırıp sana uyanı seçebilirsin.")

end_cta = render_end_cta(
    "Musashi'nin kılıcı seni bekliyor",
    "Performance ve Ultimate paketleri, sahip olduğun GeForce NOW destekli oyunları güçlü bir "
    "bilgisayar olmadan çalıştırmanı sağlar. Paketleri inceleyip kütüphanendeki oyunlarla "
    "başlayabilirsin.")


def ytembed(vid, baslik):
    return (f'<div class="gp-yt-wrap" style="max-width:720px;margin:1.6em auto;">'
            f'<iframe src="https://www.youtube.com/embed/{vid}" title="{baslik}" frameborder="0" '
            f'allow="accelerometer; autoplay; clipboard-write; encrypted-media; gyroscope; '
            f'picture-in-picture; web-share" allowfullscreen loading="lazy" '
            f'style="display:block;width:100%;aspect-ratio:16/9;height:auto;border:0;'
            f'border-radius:12px;box-shadow:0 4px 14px rgba(0,0,0,0.5);"></iframe></div>')


out, h2_no = [], 0
for tip, veri in P:
    if tip == "h1":
        out.append(f"<h1>{veri}</h1>")
    elif tip == "h2":
        h2_no += 1
        if h2_no == 2:
            out.append(cta_paketler)
        out.append(f"<h2>{veri}</h2>")
        if veri.startswith("Sıkça Sorulan"):
            out.insert(len(out) - 1, end_cta)
    elif tip == "p":
        out.append(f"<p>{veri}</p>")
        if veri.startswith("Capcom, kılıç hareketlerini"):
            out.append(editor_note)
        elif veri.startswith(f"Evet, {AD} 4 Eylül"):
            out.append(hatirlatma)
    elif tip == "yt":
        out.append(ytembed(*veri))
    elif tip == "liste":
        out.append(render_list(veri, marker="dot"))
    elif tip == "liste2":
        out.append(render_list(veri, marker="check"))
    elif tip == "puan_tablo":
        out.append(puan_tablo)

out.append(render_faq_accordion(FAQ))
out.append(render_faq_schema(FAQ))

body = "\n".join(out)
body, toc_items = inject_heading_ids(body)

rt = estimate_reading_time(body)
tldr = render_tldr([
    "<strong>Ne oynuyorsun:</strong> Genma istilasındaki Edo dönemi Kyoto'sunda Miyamoto Musashi ile "
    "geçen, savuşturma odaklı bir kılıç dövüşü aksiyonu.",
    "<strong>Seri için önemi:</strong> 2006'dan bu yana ilk ana Onimusha oyunu; remake değil, "
    "önceki oyunlardan bağımsız yeni bir hikaye.",
    "<strong>Puanlar (24 Eylül 2026):</strong> OpenCritic 86 (eleştirmenlerin %95'i tavsiye ediyor), "
    "Metacritic 85, Steam'de %90 olumlu.",
    "<strong>Süre:</strong> ana hikaye 20 saatin biraz üzerinde; yan içeriklerle 30 saat civarı.",
    "<strong>Türkçe:</strong> resmi Türkçe destek yok. Oyun çıkış günü GeForce NOW tarafında da yer "
    "aldı.",
], reading_time=rt)

info = render_info_card([
    ("Çıkış", "4 Eylül 2026"),
    ("OpenCritic", "86 Puan"),
    ("Ana Hikaye", "20 Saat+"),
    ("Tür", "Aksiyon"),
])

toc = render_floating_toc(toc_items)
m = re.search(r"</h1>", body)
body = body[:m.end()] + "\n" + toc + "\n" + tldr + "\n" + info + "\n" + body[m.end():]

body, _kat = auto_link_categories(body, max_links=2)
body, ps_bilgi = insert_preferred_source(body)
body, _link = apply_link_policy(body)
body = ensure_leading_h1(body)
final_body = wrap_gp_content(ANIMATED_BORDER_STYLE + group_into_sections(body))

print("=" * 62)
ok = print_report(verify_output(final_body, blog_type="general", expect_faq=True, yeni_icerik=True),
                  label=AD)
print("kategori linkleri:", [a for a, _ in _kat], "| tercih edilen kaynak:", ps_bilgi)

open(os.path.join(OUT, f"{SLUG}-html.txt"), "w", encoding="utf-8").write(final_body)
open(os.path.join(OUT, f"onizleme-{SLUG}.html"), "w", encoding="utf-8").write(
    embed_fonts(PAGE_HEAD.replace("__TITLE__", TITLE) + final_body + PAGE_FOOT))

from docx import Document
from docx.shared import Pt, RGBColor


def doc_renklerini_siyah_yap(doc):
    """Word'un varsayilan tema renkleri basliklari mavi basiyor; icerik kismi tamamen siyah olur."""
    for ad in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4", "List Bullet"):
        try:
            doc.styles[ad].font.color.rgb = RGBColor(0x00, 0x00, 0x00)
        except KeyError:
            pass


doc = Document()
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(11)
doc_renklerini_siyah_yap(doc)

doc.add_heading(TITLE, level=1)
for tip, veri in P:
    if tip == "h1":
        continue
    if tip == "h2":
        doc.add_heading(veri, level=2)
    elif tip == "p":
        doc.add_paragraph(re.sub(r"<[^>]+>", "", veri))
    elif tip == "yt":
        doc.add_paragraph(f"https://www.youtube.com/watch?v={veri[0]}")
    elif tip in ("liste", "liste2"):
        for md in veri:
            doc.add_paragraph(re.sub(r"<[^>]+>", "", md), style="List Bullet")
    elif tip == "puan_tablo":
        doc.add_paragraph("Puan tablosu (24 Eylül 2026): " + " | ".join(
            f"{a} {b} ({c})" for a, b, c in PUAN_SATIRLARI))
doc.add_heading("Sıkça Sorulan Sorular", level=2)
for soru, cevap in FAQ:
    doc.add_heading(soru, level=3)
    doc.add_paragraph(re.sub(r"<[^>]+>", "", cevap))

doc.add_paragraph("")
doc.add_paragraph("HTML")
for satir in final_body.split("\n"):
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(0)
    par.paragraph_format.line_spacing = 1.0
    r = par.add_run(satir)
    r.font.name = "Courier New"
    r.font.size = Pt(7)
DOC_ADI = "Onimusha Way of the Sword İncelemesi - Seri 20 Yıl Sonra Döndü"
doc.save(os.path.join(OUT, f"{DOC_ADI}.docx"))

kelime = len(re.sub(r"<[^>]+>", " ", "\n".join(v if isinstance(v, str) else "" for _, v in P)).split())
print(f"\nkelime (gövde metni): {kelime} | okuma: {rt} dk | boyut: {len(final_body)} karakter")
print("TESLİM EDİLEBİLİR" if ok else "!!! DÜZELTME GEREKLİ !!!")

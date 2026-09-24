# -*- coding: utf-8 -*-
"""GAME+ blog brief Excel'ine satir ekler ya da ayni yazinin satirini gunceller.

Yalniz YENI icerik taleplerinde kullanilir. Kullanici hazir bir taslak ilettiginde brief yazilmaz.

Sutunlar oyun detay brief'iyle (gameplus-oyun-detay-brief-icerik/scripts/brief_satiri.py) ayni
mantikta; bloga ozel olarak Icerik Tipi, Listeleme Basligi, Meta Title, Slug ve Meta Description
eklenir. Marka tablosundaki konumlara uymak icin Listeleme Basligi I, Meta Title M,
Meta Description O sutunundadir.

Kullanim:
    from blog_brief_satiri import brief_yaz
    brief_yaz("/.../Game+ Blog Briefleri.xlsx", {
        "Tarih": "2026-09-24", "Oyun / Konu": "...", "İçerik Tipi": "İnceleme", ...})

Ayni "Oyun / Konu" degerine sahip satir varsa uzerine yazilir; yoksa sona eklenir.
"""
import os

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SUTUNLAR = [
    ("Tarih", 12),                                   # A
    ("Oyun / Konu", 24),                             # B
    ("İçerik Tipi", 13),                             # C
    ("Main KW", 24),                                 # D
    ("Main KW Hacim", 14),                           # E
    ("İkincil Kelimeler", 40),                       # F
    ("Alt Başlıklar (H2)", 48),                      # G
    ("İçerik Kurgusu", 90),                          # H
    ("Listelemede gözükecek başlık (Title)", 30),    # I
    ("Link Verilecek Sayfalar", 60),                 # J
    ("SSS'ler", 56),                                 # K
    ("Yanıt Biçimi", 50),                            # L
    ("Meta Title", 30),                              # M
    ("Slug", 30),                                    # N
    ("Meta Description", 40),                        # O
    ("Durum", 14),                                   # P
]
BASLIKLAR = [ad for ad, _ in SUTUNLAR]
META_TITLE_SINIR = 60
META_DESC_SINIR = 160


def _kontrol(satir):
    eksik = [b for b in BASLIKLAR if b not in satir]
    if eksik:
        raise ValueError("brief satirinda eksik sutun: " + ", ".join(eksik))
    uyarilar = []
    if len(satir["Meta Title"]) > META_TITLE_SINIR:
        uyarilar.append(f"Meta Title {len(satir['Meta Title'])} karakter (> {META_TITLE_SINIR})")
    if len(satir["Meta Description"]) > META_DESC_SINIR:
        uyarilar.append(f"Meta Description {len(satir['Meta Description'])} karakter (> {META_DESC_SINIR})")
    if "—" in "".join(str(v) for v in satir.values()):
        raise ValueError("em dash bulundu")
    return uyarilar


def _yeni_kitap():
    wb = Workbook()
    ws = wb.active
    ws.title = "Blog Briefleri"
    for i, (ad, gen) in enumerate(SUTUNLAR, 1):
        c = ws.cell(row=1, column=i, value=ad)
        c.font = Font(name="Calibri", bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="434343")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = gen
    ws.row_dimensions[1].height = 34
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(SUTUNLAR))}1"
    return wb


def brief_yaz(yol, satir):
    uyarilar = _kontrol(satir)
    wb = load_workbook(yol) if os.path.exists(yol) else _yeni_kitap()
    ws = wb["Blog Briefleri"]
    hedef = None
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=2).value == satir["Oyun / Konu"]:
            hedef = r
            break
    if hedef is None:
        hedef = ws.max_row + 1 if ws.cell(row=ws.max_row, column=1).value else max(2, ws.max_row)
    for i, ad in enumerate(BASLIKLAR, 1):
        c = ws.cell(row=hedef, column=i, value=satir[ad])
        c.font = Font(name="Calibri", color="10332F")
        yatay = "center" if ad in ("Tarih", "İçerik Tipi", "Main KW Hacim", "Durum") else "left"
        c.alignment = Alignment(horizontal=yatay, vertical="center", wrap_text=True)
    en_uzun = max(str(satir[b]).count("\n") + 1 for b in BASLIKLAR)
    ws.row_dimensions[hedef].height = min(409, max(60, en_uzun * 15))
    wb.save(yol)
    return hedef, uyarilar
